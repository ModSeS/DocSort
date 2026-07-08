import sys
import os
import re
import io
import json
import yaml
import base64
import shutil
import tempfile
import threading
import traceback
import socket
from pathlib import Path
from datetime import datetime, timedelta
from collections import defaultdict

import requests
import fitz  # PyMuPDF
from PIL import Image

try:
    from docx import Document
    DOCX_SUPPORT = True
except ImportError:
    Document = None
    DOCX_SUPPORT = False

try:
    import olefile
    OLE_SUPPORT = True
except ImportError:
    olefile = None
    OLE_SUPPORT = False

from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QPushButton, QLabel, QLineEdit, QTextEdit, QTabWidget, QGroupBox,
    QCheckBox, QSpinBox, QFileDialog, QMessageBox, QProgressBar,
    QListWidget, QSystemTrayIcon, QMenu, QAction, QDialog,
    QTableWidget, QTableWidgetItem, QHeaderView, QComboBox,
    QDialogButtonBox, QFormLayout, QSplitter, QFrame
)
from PyQt5.QtCore import pyqtSignal, QObject, QTimer, Qt
from PyQt5.QtGui import QFont, QIcon, QColor, QBrush, QPixmap


# =========================
# Проверка единственного экземпляра
# =========================

class SingleInstance:
    """Гарантирует, что запущен только один экземпляр приложения."""

    def __init__(self, port=57342):
        self.port = port
        self.sock = None
        self.server_timer = None
        self.is_running = False

    def try_connect(self):
        """Пытается подключиться к существующему экземпляру."""
        try:
            sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
            sock.settimeout(1)
            sock.connect(('127.0.0.1', self.port))
            sock.send(b'show')
            response = sock.recv(1024)
            sock.close()
            return True
        except:
            return False

    def start_server(self, callback):
        """Запускает сервер для приема команд от новых экземпляров."""
        self.sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        self.sock.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
        self.sock.bind(('127.0.0.1', self.port))
        self.sock.listen(1)
        self.sock.setblocking(False)
        self.callback = callback
        self.is_running = True

        # Создаем таймер для проверки входящих подключений
        self.server_timer = QTimer()
        self.server_timer.timeout.connect(self.check_connections)
        self.server_timer.start(500)  # Проверка каждые 500 мс

    def check_connections(self):
        """Проверяет входящие подключения."""
        try:
            conn, addr = self.sock.accept()
            data = conn.recv(1024)
            if data == b'show':
                self.callback()
            conn.close()
        except:
            pass

    def cleanup(self):
        """Очищает ресурсы при выходе."""
        self.is_running = False
        if self.server_timer:
            self.server_timer.stop()
        if self.sock:
            self.sock.close()


# =========================
# Утилиты
# =========================

def resource_path(relative_path: str) -> str:
    """Совместимо с PyInstaller."""
    try:
        base_path = sys._MEIPASS
    except Exception:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)


def now_ts() -> str:
    return datetime.now().strftime("%H:%M:%S")


def safe_yaml_load(path: str, fallback: dict) -> dict:
    try:
        if not os.path.exists(path):
            return fallback
        with open(path, "r", encoding="utf-8") as f:
            data = yaml.safe_load(f)
        return data if isinstance(data, dict) else fallback
    except Exception:
        return fallback


def compact_spaces(text: str) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


def ensure_dir(path: str):
    Path(path).mkdir(parents=True, exist_ok=True)


def extract_json_object(text: str) -> dict:
    """Устойчивое извлечение JSON-объекта из ответа модели."""
    if not text:
        return {}

    cleaned = text.strip()
    cleaned = re.sub(r"^```(?:json)?", "", cleaned, flags=re.IGNORECASE).strip()
    cleaned = re.sub(r"```$", "", cleaned).strip()

    try:
        obj = json.loads(cleaned)
        return obj if isinstance(obj, dict) else {}
    except Exception:
        pass

    start = cleaned.find("{")
    end = cleaned.rfind("}")
    if start != -1 and end != -1 and end > start:
        candidate = cleaned[start:end + 1]
        try:
            obj = json.loads(candidate)
            return obj if isinstance(obj, dict) else {}
        except Exception:
            return {}
    return {}


def image_to_base64_png(image: Image.Image) -> str:
    buf = io.BytesIO()
    image.convert("RGB").save(buf, format="PNG")
    return base64.b64encode(buf.getvalue()).decode("ascii")


# =========================
# Диалог создания шаблона из неизвестного документа
# =========================

class CreateTemplateDialog(QDialog):
    def __init__(self, unknown_files: list, available_types: list, parent=None):
        super().__init__(parent)
        self.unknown_files = unknown_files
        self.available_types = available_types
        self.created_templates = []
        self.parent_window = parent
        self.setWindowTitle("Создание шаблонов для неизвестных документов")
        self.setModal(True)
        self.resize(1100, 650)
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout()

        title = QLabel("Найдены неизвестные документы. Создайте для них шаблоны:")
        title.setStyleSheet("font-size: 12pt; font-weight: bold; padding: 10px;")
        layout.addWidget(title)

        self.table = QTableWidget()
        self.table.setColumnCount(7)
        self.table.setHorizontalHeaderLabels([
            "Исходное имя", "Текущее имя", "Распознанный тип", "Причина",
            "Тип документа*", "Шаблон имени", "Действие"
        ])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.table.horizontalHeader().setStretchLastSection(True)

        self.table.setRowCount(len(self.unknown_files))

        for i, file_info in enumerate(self.unknown_files):
            original_name = file_info.get('original_name', file_info.get('filename', ''))
            original_item = QTableWidgetItem(original_name)
            original_item.setBackground(QColor(240, 248, 255))
            self.table.setItem(i, 0, original_item)

            current_name = file_info.get('filename', '')
            current_item = QTableWidgetItem(current_name)
            current_item.setBackground(QColor(255, 255, 240))
            self.table.setItem(i, 1, current_item)

            type_item = QTableWidgetItem(file_info.get('type', 'Неизвестно'))
            type_item.setBackground(QColor(255, 255, 200))
            self.table.setItem(i, 2, type_item)

            reason_item = QTableWidgetItem(file_info.get('reason', ''))
            reason_item.setBackground(QColor(255, 240, 240))
            self.table.setItem(i, 3, reason_item)

            type_combo = QComboBox()
            type_combo.setEditable(True)
            type_combo.addItems(self.available_types)
            recognized_type = file_info.get('type', '')
            if recognized_type and recognized_type not in self.available_types:
                type_combo.insertItem(0, recognized_type)
                type_combo.setCurrentIndex(0)
            elif recognized_type in self.available_types:
                type_combo.setCurrentText(recognized_type)
            type_combo.currentTextChanged.connect(lambda text, row=i: self.on_type_changed(row, text))
            self.table.setCellWidget(i, 4, type_combo)

            filename_edit = QLineEdit()
            filename_edit.setPlaceholderText("{тип} №{номер} от {дата}")
            filename_edit.setText("{тип} №{номер} от {дата}")
            self.table.setCellWidget(i, 5, filename_edit)

            apply_btn = QPushButton("Применить сейчас")
            apply_btn.setStyleSheet("""
                QPushButton {
                    background-color: #4CAF50;
                    color: white;
                    border: none;
                    padding: 5px 10px;
                    border-radius: 3px;
                }
                QPushButton:hover {
                    background-color: #45a049;
                }
            """)
            apply_btn.clicked.connect(lambda checked, row=i: self.apply_existing_template(row))
            self.table.setCellWidget(i, 6, apply_btn)

        self.table.setColumnWidth(0, 180)
        self.table.setColumnWidth(1, 180)
        self.table.setColumnWidth(2, 120)
        self.table.setColumnWidth(3, 180)
        self.table.setColumnWidth(4, 130)
        self.table.setColumnWidth(5, 180)

        layout.addWidget(self.table)

        self.remaining_label = QLabel(f"Осталось обработать: {len(self.unknown_files)}")
        self.remaining_label.setStyleSheet("font-size: 11pt; padding: 5px; color: #555;")
        layout.addWidget(self.remaining_label)

        apply_all_layout = QHBoxLayout()
        apply_all_btn = QPushButton("Применить шаблон ко всем оставшимся")
        apply_all_btn.clicked.connect(self.apply_to_all)
        apply_all_layout.addWidget(apply_all_btn)

        apply_all_layout.addWidget(QLabel("Шаблон:"))
        self.global_template_edit = QLineEdit("{тип} №{номер} от {дата}")
        apply_all_layout.addWidget(self.global_template_edit)

        layout.addLayout(apply_all_layout)

        hint = QLabel(
            "Подсказка: Выберите тип документа из списка или введите новый.\n"
            "Шаблон имени может содержать переменные: {тип}, {номер}, {дата}, {тема}\n"
            "Кнопка 'Применить сейчас' сразу сортирует файл и убирает его из списка."
        )
        hint.setStyleSheet("color: #666; font-size: 10pt; padding: 5px;")
        layout.addWidget(hint)

        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.accepted.connect(self.accept_templates)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

        self.setLayout(layout)

    def get_config(self) -> dict:
        if self.parent_window:
            if hasattr(self.parent_window, 'parent_window') and self.parent_window.parent_window:
                return self.parent_window.parent_window.config
            elif hasattr(self.parent_window, 'config'):
                return self.parent_window.config
        return {"out_dir": "sorted", "unknown_dir": "unknown"}

    def on_type_changed(self, row: int, text: str):
        if not self.parent_window:
            return

        templates = self.parent_window.templates.get("templates", [])
        for tpl in templates:
            if tpl.get("doc_type", "").lower() == text.lower():
                filename_edit = self.table.cellWidget(row, 5)
                if filename_edit:
                    filename_edit.setText(tpl.get("filename", "{тип} №{номер} от {дата}"))
                break

    def apply_existing_template(self, row: int):
        type_combo = self.table.cellWidget(row, 4)
        filename_edit = self.table.cellWidget(row, 5)

        if not type_combo or not filename_edit:
            return

        doc_type = type_combo.currentText().strip()
        filename_template = filename_edit.text().strip()

        if not doc_type:
            QMessageBox.warning(self, "Внимание", "Выберите тип документа")
            return

        if self.parent_window:
            templates = self.parent_window.templates.get("templates", [])
            existing_template = None
            for tpl in templates:
                if tpl.get("doc_type", "").lower() == doc_type.lower():
                    existing_template = tpl
                    break

            if existing_template:
                file_info = self.unknown_files[row]
                src_path = file_info.get('filepath', '')
                if not src_path or not os.path.exists(src_path):
                    QMessageBox.warning(self, "Ошибка", "Исходный файл не найден")
                    return

                try:
                    self.quick_sort_file(src_path, existing_template, file_info)

                    clean_template = {
                        "doc_type": existing_template.get("doc_type", doc_type),
                        "filename": existing_template.get("filename", filename_template),
                        "subdir": existing_template.get("subdir", doc_type)
                    }
                    if clean_template not in self.created_templates:
                        self.created_templates.append(clean_template)

                    self.remove_row(row)

                    QMessageBox.information(
                        self,
                        "Успех",
                        f"Файл '{file_info.get('original_name', file_info.get('filename', ''))}' успешно отсортирован как '{doc_type}'"
                    )

                except Exception as e:
                    QMessageBox.critical(self, "Ошибка", f"Не удалось обработать файл: {e}")
                    traceback.print_exc()
            else:
                new_template = {
                    "doc_type": doc_type,
                    "filename": filename_template,
                    "subdir": doc_type
                }

                if new_template not in self.created_templates:
                    self.created_templates.append(new_template)

                file_info = self.unknown_files[row]
                src_path = file_info.get('filepath', '')

                if src_path and os.path.exists(src_path):
                    try:
                        self.quick_sort_file(src_path, new_template, file_info)
                        self.remove_row(row)

                    except Exception as e:
                        QMessageBox.critical(self, "Ошибка", f"Не удалось обработать файл: {e}")
                        traceback.print_exc()
                else:
                    QMessageBox.warning(self, "Ошибка", "Исходный файл не найден")

    def remove_row(self, row: int):
        if 0 <= row < len(self.unknown_files):
            del self.unknown_files[row]

        self.table.removeRow(row)
        self.update_row_handlers()

        remaining = self.table.rowCount()
        self.remaining_label.setText(f"Осталось обработать: {remaining}")

    def update_row_handlers(self):
        for i in range(self.table.rowCount()):
            type_combo = self.table.cellWidget(i, 4)
            if type_combo:
                try:
                    type_combo.currentTextChanged.disconnect()
                except:
                    pass
                type_combo.currentTextChanged.connect(lambda text, row=i: self.on_type_changed(row, text))

            apply_btn = self.table.cellWidget(i, 6)
            if apply_btn:
                try:
                    apply_btn.clicked.disconnect()
                except:
                    pass
                apply_btn.clicked.connect(lambda checked, row=i: self.apply_existing_template(row))

    def quick_sort_file(self, src_path: str, template: dict, file_info: dict):
        meta = file_info.get('meta', {})

        if not meta:
            meta = {
                "type": file_info.get('type', 'Неизвестный'),
                "number": "",
                "date": "",
                "designation": "Без темы"
            }

        number = compact_spaces(meta.get("number", ""))
        date_iso = compact_spaces(meta.get("date", ""))
        designation = compact_spaces(meta.get("designation", "Без темы"))

        date_formatted = ""
        if date_iso:
            try:
                date_obj = datetime.strptime(date_iso, "%Y-%m-%d")
                date_formatted = date_obj.strftime("%d.%m.%Y")
            except:
                pass

        values = {
            "тип": template.get("doc_type", "Документ"),
            "номер": number if number else "без номера",
            "дата": date_formatted if date_formatted else "без даты",
            "тема": designation if designation != "Без темы" else template.get("doc_type", "Документ"),
        }

        filename_tpl = template.get("filename", "{тип} №{номер} от {дата}")
        try:
            name = filename_tpl.format(**values)
        except Exception:
            name = "{тип} №{номер} от {дата}".format(**values)

        name = re.sub(r"№\s*№+", "№", name)
        INVALID_CHARS = r'<>:"/\|?*'
        for ch in INVALID_CHARS:
            name = name.replace(ch, "_")
        name = re.sub(r"[\r\n\t]+", " ", name)
        name = re.sub(r"\s+", " ", name).strip()
        name = name.rstrip(". ")
        if len(name) > 180:
            name = name[:180].rstrip(". ")
        name = name or "Без названия"

        ext = Path(src_path).suffix.lower()
        subdir = template.get("subdir", template.get("doc_type", ""))

        for ch in INVALID_CHARS:
            subdir = subdir.replace(ch, "_")
        subdir = re.sub(r"[\r\n\t]+", " ", subdir)
        subdir = re.sub(r"\s+", " ", subdir).strip()

        config = self.get_config()
        out_dir = config.get("out_dir", "sorted")
        dest_dir = os.path.join(out_dir, subdir) if subdir else out_dir
        ensure_dir(dest_dir)

        dest_path = os.path.join(dest_dir, name + ext)

        counter = 2
        while os.path.exists(dest_path):
            name_with_counter = f"{name} ({counter})"
            dest_path = os.path.join(dest_dir, name_with_counter + ext)
            counter += 1

        shutil.move(src_path, dest_path)

        ensure_dir("logs")
        history_entry = {
            "src": str(src_path),
            "dest": str(dest_path),
            "status": "ok_manual",
            "meta": meta,
            "template": template,
            "timestamp": datetime.now().isoformat(),
            "method": "quick_sort"
        }

        with open("logs/history.jsonl", "a", encoding="utf-8") as f:
            f.write(json.dumps(history_entry, ensure_ascii=False) + "\n")

        return dest_path

    def apply_to_all(self):
        template = self.global_template_edit.text()
        if not template:
            QMessageBox.warning(self, "Внимание", "Введите шаблон имени")
            return

        for i in range(self.table.rowCount()):
            filename_edit = self.table.cellWidget(i, 5)
            if filename_edit:
                filename_edit.setText(template)

        QMessageBox.information(
            self,
            "Применено",
            f"Шаблон применён к {self.table.rowCount()} оставшимся файлам.\n"
            "Теперь нажмите 'Применить сейчас' для каждого файла или OK для сохранения шаблонов."
        )

    def accept_templates(self):
        for i in range(self.table.rowCount()):
            type_combo = self.table.cellWidget(i, 4)
            filename_edit = self.table.cellWidget(i, 5)

            doc_type = type_combo.currentText().strip() if type_combo else ""
            filename = filename_edit.text().strip() if filename_edit else ""

            if doc_type and filename:
                template = {
                    "doc_type": doc_type,
                    "filename": filename,
                    "subdir": doc_type
                }

                if template not in self.created_templates:
                    self.created_templates.append(template)

                    file_info = self.unknown_files[i] if i < len(self.unknown_files) else {}
                    template['source_file'] = file_info.get('filepath', '')
                    template['original_type'] = file_info.get('type', '')

        if not self.created_templates:
            QMessageBox.warning(self, "Внимание", "Не создано ни одного шаблона")
            return

        for i in range(self.table.rowCount()):
            type_combo = self.table.cellWidget(i, 4)
            doc_type = type_combo.currentText().strip() if type_combo else ""

            for template in self.created_templates:
                if template["doc_type"].lower() == doc_type.lower():
                    if i < len(self.unknown_files):
                        file_info = self.unknown_files[i]
                        src_path = file_info.get('filepath', '')

                        if src_path and os.path.exists(src_path):
                            try:
                                self.quick_sort_file(src_path, template, file_info)
                            except Exception as e:
                                print(f"Ошибка при сортировке файла {src_path}: {e}")
                                traceback.print_exc()
                    break

        self.accept()


# =========================
# Окно уведомлений о результатах
# =========================

class ResultsDialog(QDialog):
    BAD_TYPES = {"утверждение", "соглашение", "утверждаю", "согласовано",
                 "подпись", "подписи", "печать", "штамп"}

    def __init__(self, results: dict, templates: dict, inbox_dir: str, parent=None):
        super().__init__(parent)
        self.results = results
        self.templates = templates
        self.inbox_dir = inbox_dir
        self.parent_window = parent
        self.config = self.get_config_from_parent()
        self.setWindowTitle("Результаты обработки документов")
        self.setModal(True)
        self.resize(1050, 650)
        self.init_ui()

    def get_config_from_parent(self) -> dict:
        if self.parent_window and hasattr(self.parent_window, 'config'):
            return self.parent_window.config
        return {"out_dir": "sorted", "unknown_dir": "unknown", "inbox_dir": "scannerdata"}

    def init_ui(self):
        layout = QVBoxLayout()

        title = QLabel(f"Обработка завершена: {datetime.now().strftime('%d.%m.%Y %H:%M:%S')}")
        title.setStyleSheet("font-size: 14pt; font-weight: bold; padding: 10px;")
        layout.addWidget(title)

        stats_layout = QHBoxLayout()

        self.processed_label = QLabel(f"✅ Успешно: {self.results.get('processed', 0)}")
        self.processed_label.setStyleSheet("font-size: 12pt; padding: 10px; color: green;")
        stats_layout.addWidget(self.processed_label)

        self.unknown_label = QLabel(f"❓ Неизвестные: {self.results.get('unknown', 0)}")
        self.unknown_label.setStyleSheet("font-size: 12pt; padding: 10px; color: red;")
        stats_layout.addWidget(self.unknown_label)

        total_label = QLabel(f"📁 Всего: {self.results.get('total', 0)}")
        total_label.setStyleSheet("font-size: 12pt; padding: 10px;")
        stats_layout.addWidget(total_label)

        stats_layout.addStretch()
        layout.addLayout(stats_layout)

        self.details = self.results.get('details', [])
        if self.details:
            self.table = QTableWidget()
            self.table.setColumnCount(7)
            self.table.setHorizontalHeaderLabels([
                "Исходное имя", "Файл", "Тип", "Статус",
                "Причина", "Путь", "Действие"
            ])
            self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
            self.table.horizontalHeader().setStretchLastSection(True)

            self.refresh_table()

            self.table.setColumnWidth(0, 200)
            self.table.setColumnWidth(1, 200)
            self.table.setColumnWidth(2, 120)
            self.table.setColumnWidth(3, 80)
            self.table.setColumnWidth(4, 180)
            self.table.setColumnWidth(5, 200)
            self.table.setColumnWidth(6, 100)

            layout.addWidget(self.table)

        buttons_layout = QHBoxLayout()

        report_btn = QPushButton("📊 Создать отчёт")
        report_btn.setStyleSheet("""
            QPushButton {
                font-size: 11pt; 
                padding: 10px;
                background-color: #9C27B0;
                color: white;
                border: none;
                border-radius: 5px;
            }
            QPushButton:hover {
                background-color: #7B1FA2;
            }
        """)
        report_btn.clicked.connect(self.generate_report)
        buttons_layout.addWidget(report_btn)

        if self.results.get('unknown', 0) > 0:
            create_templates_btn = QPushButton("📝 Создать шаблоны")
            create_templates_btn.setStyleSheet("""
                QPushButton {
                    font-size: 11pt; 
                    padding: 10px;
                    background-color: #2196F3;
                    color: white;
                    border: none;
                    border-radius: 5px;
                }
                QPushButton:hover {
                    background-color: #1976D2;
                }
            """)
            create_templates_btn.clicked.connect(self.create_templates_for_unknown)
            buttons_layout.addWidget(create_templates_btn)

        open_results_btn = QPushButton("📂 Результаты")
        open_results_btn.setStyleSheet("""
            QPushButton {
                font-size: 11pt; 
                padding: 10px;
                background-color: #607D8B;
                color: white;
                border: none;
                border-radius: 5px;
            }
            QPushButton:hover {
                background-color: #455A64;
            }
        """)
        open_results_btn.clicked.connect(self.open_results_folder)
        buttons_layout.addWidget(open_results_btn)

        buttons_layout.addStretch()

        close_btn = QPushButton("Закрыть")
        close_btn.clicked.connect(self.accept)
        buttons_layout.addWidget(close_btn)

        layout.addLayout(buttons_layout)

        self.setLayout(layout)

    def refresh_table(self):
        self.table.setRowCount(len(self.details))

        for i, detail in enumerate(self.details):
            original_name = detail.get('original_name', detail.get('filename', ''))
            self.table.setItem(i, 0, QTableWidgetItem(original_name))

            new_filename = detail.get('filename', '')
            self.table.setItem(i, 1, QTableWidgetItem(new_filename))

            self.table.setItem(i, 2, QTableWidgetItem(detail.get('type', '')))

            status_item = QTableWidgetItem(detail.get('status', ''))
            if detail.get('status') == 'OK':
                status_item.setBackground(QColor(200, 255, 200))
            else:
                status_item.setBackground(QColor(255, 200, 200))
            self.table.setItem(i, 3, status_item)

            self.table.setItem(i, 4, QTableWidgetItem(detail.get('reason', '')))
            self.table.setItem(i, 5, QTableWidgetItem(detail.get('destination', '')))

            if detail.get('status') == 'OK':
                open_btn = QPushButton("📂 Открыть")
                open_btn.setStyleSheet("""
                    QPushButton {
                        background-color: #4CAF50;
                        color: white;
                        border: none;
                        padding: 3px 8px;
                        border-radius: 3px;
                        font-size: 9pt;
                    }
                    QPushButton:hover {
                        background-color: #45a049;
                    }
                """)
                open_btn.clicked.connect(
                    lambda checked, dest=detail.get('destination', ''): self.open_folder(dest)
                )
                self.table.setCellWidget(i, 6, open_btn)
            else:
                return_btn = QPushButton("↩️ Вернуть")
                return_btn.setStyleSheet("""
                    QPushButton {
                        background-color: #FF9800;
                        color: white;
                        border: none;
                        padding: 3px 8px;
                        border-radius: 3px;
                        font-size: 9pt;
                    }
                    QPushButton:hover {
                        background-color: #F57C00;
                    }
                """)
                return_btn.clicked.connect(
                    lambda checked, row=i, detail=detail: self.return_to_inbox(row, detail)
                )
                self.table.setCellWidget(i, 6, return_btn)

    def update_stats(self):
        processed = sum(1 for d in self.details if d.get('status') == 'OK')
        unknown = sum(1 for d in self.details if d.get('status') != 'OK')
        total = len(self.details)

        self.processed_label.setText(f"✅ Успешно: {processed}")
        self.unknown_label.setText(f"❓ Неизвестные: {unknown}")

        self.results['processed'] = processed
        self.results['unknown'] = unknown
        self.results['total'] = total

    def return_to_inbox(self, row: int, detail: dict):
        original_name = detail.get('original_name', detail.get('filename', ''))
        current_filename = detail.get('filename', '')
        unknown_dir = self.config.get("unknown_dir", "unknown")
        inbox_dir = self.config.get("inbox_dir", "scannerdata")

        src_path = os.path.join(unknown_dir, current_filename)

        if not os.path.exists(src_path):
            QMessageBox.warning(self, "Ошибка", f"Файл не найден:\n{src_path}")
            return

        dest_path = os.path.join(inbox_dir, original_name)

        if os.path.exists(dest_path):
            base, ext = os.path.splitext(original_name)
            counter = 2
            while os.path.exists(os.path.join(inbox_dir, f"{base} ({counter}){ext}")):
                counter += 1
            new_original_name = f"{base} ({counter}){ext}"
            dest_path = os.path.join(inbox_dir, new_original_name)
            returned_name = new_original_name
        else:
            returned_name = original_name

        try:
            shutil.move(src_path, dest_path)

            del self.details[row]
            self.refresh_table()
            self.update_stats()
            self.log_return(original_name, returned_name)

            QMessageBox.information(
                self,
                "Успех",
                f"Файл возвращён в папку распознавания:\n{returned_name}"
            )

            if self.results.get('unknown', 0) == 0:
                self.update_buttons_state()

        except Exception as e:
            QMessageBox.critical(self, "Ошибка", f"Не удалось переместить файл:\n{e}")
            traceback.print_exc()

    def log_return(self, original_name: str, returned_name: str):
        try:
            ensure_dir("logs")
            log_entry = {
                "action": "return_to_inbox",
                "original_name": original_name,
                "returned_name": returned_name,
                "timestamp": datetime.now().isoformat()
            }
            with open("logs/returns.jsonl", "a", encoding="utf-8") as f:
                f.write(json.dumps(log_entry, ensure_ascii=False) + "\n")
        except Exception:
            pass

    def update_buttons_state(self):
        for child in self.findChildren(QPushButton):
            if child.text() == "📝 Создать шаблоны":
                child.setVisible(self.results.get('unknown', 0) > 0)
                break

    def generate_report(self):
        try:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            report_path = f"logs/report_{timestamp}.html"

            total = self.results.get('total', 0)
            processed = self.results.get('processed', 0)
            unknown = self.results.get('unknown', 0)
            details = self.details

            html_content = f"""<!DOCTYPE html>
<html lang="ru">
<head>
    <meta charset="UTF-8">
    <title>Отчёт обработки документов</title>
    <style>
        body {{ 
            font-family: Arial, sans-serif; 
            margin: 20px; 
            background: #f5f5f5; 
        }}
        .container {{ 
            max-width: 1000px; 
            margin: 0 auto; 
            background: white; 
            padding: 20px; 
            border-radius: 8px; 
            box-shadow: 0 2px 10px rgba(0,0,0,0.1); 
        }}
        h1 {{ color: #333; border-bottom: 2px solid #4CAF50; padding-bottom: 10px; }}
        h2 {{ color: #555; margin-top: 30px; }}
        .stats {{ 
            display: flex; 
            gap: 20px; 
            margin: 20px 0; 
        }}
        .stat {{ 
            flex: 1; 
            padding: 15px; 
            background: #f9f9f9; 
            border-radius: 5px; 
            text-align: center; 
        }}
        .stat .num {{ font-size: 24px; font-weight: bold; }}
        .success {{ color: #4CAF50; }}
        .danger {{ color: #f44336; }}
        .neutral {{ color: #2196F3; }}
        table {{ 
            width: 100%; 
            border-collapse: collapse; 
            margin: 15px 0; 
        }}
        th {{ 
            background: #4CAF50; 
            color: white; 
            padding: 10px; 
            text-align: left; 
        }}
        td {{ 
            padding: 8px 10px; 
            border-bottom: 1px solid #ddd; 
        }}
        tr:hover {{ background: #f5f5f5; }}
        .badge {{
            padding: 3px 8px;
            border-radius: 3px;
            font-size: 12px;
        }}
        .badge-ok {{ background: #d4edda; color: #155724; }}
        .badge-unknown {{ background: #f8d7da; color: #721c24; }}
        .footer {{ 
            margin-top: 30px; 
            padding-top: 15px; 
            border-top: 1px solid #ddd; 
            color: #999; 
            font-size: 12px; 
        }}
    </style>
</head>
<body>
    <div class="container">
        <h1>📊 Отчёт обработки документов</h1>
        <p>Создан: {datetime.now().strftime('%d.%m.%Y %H:%M:%S')}</p>
        
        <div class="stats">
            <div class="stat">
                <div class="num neutral">{total}</div>
                <div>Всего файлов</div>
            </div>
            <div class="stat">
                <div class="num success">{processed}</div>
                <div>Успешно</div>
            </div>
            <div class="stat">
                <div class="num danger">{unknown}</div>
                <div>Неизвестные</div>
            </div>
            <div class="stat">
                <div class="num">{int(processed/total*100) if total > 0 else 0}%</div>
                <div>Успешность</div>
            </div>
        </div>
        
        <h2>✅ Успешно обработанные ({processed})</h2>
        <table>
            <tr>
                <th>Исходное имя</th>
                <th>Новое имя</th>
                <th>Тип</th>
                <th>Путь</th>
            </tr>
"""

            success_details = [d for d in details if d.get('status') == 'OK']
            for detail in success_details:
                original_name = detail.get('original_name', detail.get('filename', ''))
                new_name = detail.get('filename', '')
                doc_type = detail.get('type', '')
                destination = detail.get('destination', '')

                html_content += f"""
            <tr>
                <td>{original_name}</td>
                <td>{new_name}</td>
                <td><span class="badge badge-ok">{doc_type}</span></td>
                <td>{destination}</td>
            </tr>
"""

            if not success_details:
                html_content += """
            <tr><td colspan="4" style="text-align: center; color: #999;">Нет успешно обработанных файлов</td></tr>
"""

            html_content += f"""
        </table>
        
        <h2>❓ Неизвестные ({unknown})</h2>
        <table>
            <tr>
                <th>Исходное имя</th>
                <th>Текущее имя</th>
                <th>Тип</th>
                <th>Причина</th>
            </tr>
"""

            unknown_details = [d for d in details if d.get('status') != 'OK']
            for detail in unknown_details:
                original_name = detail.get('original_name', detail.get('filename', ''))
                current_name = detail.get('filename', '')
                doc_type = detail.get('type', 'Неизвестный')
                reason = detail.get('reason', '')

                html_content += f"""
            <tr>
                <td>{original_name}</td>
                <td>{current_name}</td>
                <td><span class="badge badge-unknown">{doc_type}</span></td>
                <td>{reason}</td>
            </tr>
"""

            if not unknown_details:
                html_content += """
            <tr><td colspan="4" style="text-align: center; color: #999;">Нет неизвестных файлов</td></tr>
"""

            html_content += f"""
        </table>
        
        <div class="footer">
            ДокСорт v2.3 | {datetime.now().strftime('%d.%m.%Y')}
        </div>
    </div>
</body>
</html>"""

            ensure_dir("logs")
            with open(report_path, "w", encoding="utf-8") as f:
                f.write(html_content)

            import webbrowser
            webbrowser.open(f"file://{os.path.abspath(report_path)}")

            QMessageBox.information(self, "Отчёт создан", f"Отчёт сохранён:\n{report_path}")

        except Exception as e:
            QMessageBox.critical(self, "Ошибка", f"Не удалось создать отчёт:\n{e}")
            traceback.print_exc()

    def open_folder(self, dest_path: str):
        import subprocess

        if not dest_path:
            return

        full_path = os.path.abspath(dest_path)

        if not os.path.exists(full_path):
            QMessageBox.warning(self, "Ошибка", f"Файл не найден:\n{full_path}")
            return

        try:
            if sys.platform == 'win32':
                subprocess.run(['explorer', '/select,', full_path])
            elif sys.platform == 'darwin':
                subprocess.run(['open', '-R', full_path])
            else:
                folder_path = os.path.dirname(full_path)
                subprocess.run(['xdg-open', folder_path])
        except Exception as e:
            QMessageBox.warning(self, "Ошибка", f"Не удалось открыть папку:\n{e}")

    def open_results_folder(self):
        import subprocess

        out_dir = self.config.get("out_dir", "sorted")
        full_path = os.path.abspath(out_dir)

        if not os.path.exists(full_path):
            QMessageBox.warning(self, "Ошибка", f"Папка не найдена:\n{full_path}")
            return

        try:
            if sys.platform == 'win32':
                os.startfile(full_path)
            elif sys.platform == 'darwin':
                subprocess.run(['open', full_path])
            else:
                subprocess.run(['xdg-open', full_path])
        except Exception as e:
            QMessageBox.warning(self, "Ошибка", f"Не удалось открыть папку:\n{e}")

    def get_available_types(self) -> list:
        types = []
        for tpl in self.templates.get("templates", []):
            doc_type = tpl.get("doc_type", "")
            if doc_type and doc_type not in types:
                types.append(doc_type)
        return sorted(types)

    def get_unknown_files_info(self) -> list:
        unknown_files = []
        unknown_dir = self.config.get("unknown_dir", "unknown")

        for detail in self.details:
            if detail.get('status') != 'OK':
                filename = detail.get('filename', '')
                original_name = detail.get('original_name', filename)
                filepath = os.path.join(unknown_dir, filename)

                meta = {}
                try:
                    history_path = "logs/history.jsonl"
                    if os.path.exists(history_path):
                        with open(history_path, "r", encoding="utf-8") as f:
                            for line in f:
                                try:
                                    entry = json.loads(line)
                                    if entry.get("status") == "unknown" and Path(entry.get("src", "")).name == detail.get('filename', ''):
                                        meta = entry.get("meta", {})
                                        break
                                except:
                                    pass
                except:
                    pass

                file_info = {
                    'filename': detail.get('filename', ''),
                    'original_name': original_name,
                    'type': detail.get('type', 'Неизвестно'),
                    'reason': detail.get('reason', ''),
                    'filepath': filepath,
                    'meta': meta
                }
                unknown_files.append(file_info)
        return unknown_files

    def create_templates_for_unknown(self):
        unknown_files = self.get_unknown_files_info()
        if not unknown_files:
            QMessageBox.information(self, "Информация", "Нет неизвестных файлов")
            return

        available_types = self.get_available_types()
        dialog = CreateTemplateDialog(unknown_files, available_types, self)

        if dialog.exec_() == QDialog.Accepted and dialog.created_templates:
            current_templates = self.templates.get("templates", [])

            for template in dialog.created_templates:
                clean_template = {
                    "doc_type": template["doc_type"],
                    "filename": template["filename"],
                    "subdir": template["subdir"]
                }

                if clean_template not in current_templates:
                    current_templates.append(clean_template)

            self.templates["templates"] = current_templates

            with open("templates.yaml", "w", encoding="utf-8") as f:
                yaml.dump(self.templates, f, allow_unicode=True, default_flow_style=False, sort_keys=False)

            if self.parent_window:
                self.parent_window.templates = self.templates
                self.parent_window.templates_widget.templates = current_templates
                self.parent_window.templates_widget.refresh_list()

            QMessageBox.information(
                self,
                "Готово",
                f"Создано шаблонов: {len(dialog.created_templates)}\n"
                f"Всего шаблонов: {len(current_templates)}\n\n"
                "Файлы были перемещены согласно новым шаблонам."
            )


# =========================
# Signals
# =========================

class WorkerSignals(QObject):
    finished = pyqtSignal(dict)
    error = pyqtSignal(str)
    progress = pyqtSignal(int, str)
    log = pyqtSignal(str)


# =========================
# Главный worker
# =========================

class DocumentProcessorWorker(threading.Thread):
    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc", ".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp"}
    INVALID_CHARS = r'<>:"/\|?*'

    def __init__(self, config: dict, templates: dict, inbox_dir: str):
        super().__init__()
        self.config = config
        self.templates = templates
        self.inbox_dir = inbox_dir
        self.signals = WorkerSignals()
        self.is_running = True
        self.results = {
            'processed': 0,
            'unknown': 0,
            'total': 0,
            'details': []
        }

    def log(self, msg: str):
        self.signals.log.emit(msg)

    def sanitize_filename(self, name: str) -> str:
        name = str(name or "")
        for ch in self.INVALID_CHARS:
            name = name.replace(ch, "_")
        name = re.sub(r"[\r\n\t]+", " ", name)
        name = re.sub(r"\s+", " ", name).strip()
        name = name.rstrip(". ")
        if len(name) > 180:
            name = name[:180].rstrip(". ")
        return name or "Без названия"

    def get_unique_path(self, path: str) -> str:
        base, ext = os.path.splitext(path)
        if not os.path.exists(path):
            return path
        i = 2
        while True:
            candidate = f"{base} ({i}){ext}"
            if not os.path.exists(candidate):
                return candidate
            i += 1

    def available_doc_types(self) -> list[str]:
        out = []
        for tpl in self.templates.get("templates", []):
            doc_type = compact_spaces(tpl.get("doc_type", ""))
            if doc_type and doc_type not in out:
                out.append(doc_type)
        return out

    def write_history(self, obj: dict):
        ensure_dir("logs")
        with open("logs/history.jsonl", "a", encoding="utf-8") as f:
            f.write(json.dumps(obj, ensure_ascii=False) + "\n")

    def extract_text_from_docx(self, file_path: str) -> str:
        if not DOCX_SUPPORT:
            self.log("python-docx не установлен")
            return ""
        try:
            doc = Document(file_path)
            parts = []
            for p in doc.paragraphs:
                if p.text.strip():
                    parts.append(p.text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        cell_text = compact_spaces(cell.text)
                        if cell_text:
                            row_text.append(cell_text)
                    if row_text:
                        parts.append(" | ".join(row_text))
            return "\n".join(parts).strip()
        except Exception as e:
            self.log(f"Ошибка чтения DOCX: {e}")
            return ""

    def extract_text_from_doc(self, file_path: str) -> str:
        if not OLE_SUPPORT:
            self.log("olefile не установлен: .doc поддерживается ограниченно")
        try:
            if OLE_SUPPORT and olefile.isOleFile(file_path):
                ole = olefile.OleFileIO(file_path)
                if ole.exists("WordDocument"):
                    data = ole.openstream("WordDocument").read()
                    ole.close()
                    for enc in ("cp1251", "utf-16le", "cp1252", "latin-1"):
                        txt = data.decode(enc, errors="ignore")
                        txt = re.sub(r"[^\w\s\.,;:!?№/\-()«»А-Яа-яЁё]", " ", txt)
                        txt = compact_spaces(txt)
                        if len(txt) > 80:
                            return txt
        except Exception as e:
            self.log(f"Ошибка чтения DOC: {e}")

        try:
            data = Path(file_path).read_bytes()
            best = ""
            for enc in ("cp1251", "utf-16le", "latin-1"):
                txt = data.decode(enc, errors="ignore")
                txt = re.sub(r"[^\w\s\.,;:!?№/\-()«»А-Яа-яЁё]", " ", txt)
                txt = compact_spaces(txt)
                if len(txt) > len(best):
                    best = txt
            return best
        except Exception:
            return ""

    def render_pdf_pages_to_images(self, pdf_path: str, max_pages: int, dpi: int) -> list[Image.Image]:
        images = []
        try:
            doc = fitz.open(pdf_path)
            pages = min(len(doc), max_pages)
            zoom = dpi / 72.0
            mat = fitz.Matrix(zoom, zoom)
            for i in range(pages):
                page = doc[i]
                pix = page.get_pixmap(matrix=mat, alpha=False)
                img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                images.append(img)
            doc.close()
        except Exception as e:
            self.log(f"Ошибка рендера PDF: {e}")
        return images

    def build_header_crop(self, img: Image.Image, ratio: float = 0.35) -> Image.Image:
        w, h = img.size
        crop_h = max(1, int(h * ratio))
        return img.crop((0, 0, w, crop_h))

    def build_vision_prompt(self) -> str:
        user_hint = compact_spaces(self.config.get("vision_user_prompt", ""))

        prompt = f"""
Ты — система анализа русских документов по изображению.

Тебе нужно определить и вернуть ТОЛЬКО JSON:
{{
  "type": "",
  "number": "",
  "date": "",
  "designation": "",
  "raw_text": ""
}}

ПРАВИЛА:
1. type:
   - Определи тип документа (например: Решение, Приказ, Договор, Акт, Заявление, Счет, Накладная и т.д.).
   - Верни тип документа в именительном падеже с заглавной буквы.
   - Если не уверен — верни "".

2. number:
   - Извлекай только номер документа.
   - Не включай символ "№".
   - Не включай слова типа "Решение", "Приказ", "Договор" и эрративы слов с цифрой внутри.
   - Если номера нет или в нём нет цифр — верни "".

3. date:
   - Верни дату документа строго в формате YYYY-MM-DD.
   - Бери дату документа, а не случайную дату внутри текста.
   - Если даты нет — верни "".

4. designation:
   - Краткая тема документа 3-7 слов.
   - Если непонятно — "Без темы".

5. raw_text:
   - Верни как можно больше распознанного текста.

ВАЖНО:
- Ответ должен быть СТРОГО JSON.
- Не используй markdown.
- Не добавляй пояснений.

Дополнительные указания:
{user_hint if user_hint else "Нет дополнительных указаний."}
""".strip()
        return prompt

    def call_vision_model(self, image: Image.Image) -> dict:
        if not self.config.get("vision_enabled", True):
            return {}

        url = self.config.get("vision_url", "http://localhost:11434/api/generate").strip()
        model = self.config.get("vision_model", "qwen2.5vl:7b").strip()
        timeout = int(self.config.get("vision_timeout", 0))
        prompt = self.build_vision_prompt()
        b64 = image_to_base64_png(image)

        payload = {
            "model": model,
            "prompt": prompt,
            "images": [b64],
            "stream": False,
            "format": "json",
            "options": {
                "temperature": 0,
                "num_ctx": int(self.config.get("vision_num_ctx", 8192))
            }
        }

        ensure_dir("logs")
        Path("logs/last_vision_prompt.txt").write_text(prompt, encoding="utf-8")

        if timeout and timeout > 0:
            r = requests.post(url, json=payload, timeout=timeout)
        else:
            r = requests.post(url, json=payload)
        r.raise_for_status()
        data = r.json()

        content = (data.get("response") or "").strip()
        Path("logs/last_vision_output.txt").write_text(content, encoding="utf-8")

        if not content:
            self.log("Vision-модель вернула пустой response")
            return {}

        obj = extract_json_object(content)
        if not obj:
            self.log(f"Vision-модель не вернула JSON. Первые 300 символов: {content[:300]}")
            return {}

        return obj

    def call_vision_for_file(self, file_path: str) -> dict:
        ext = Path(file_path).suffix.lower()
        max_pages = int(self.config.get("vision_max_pages", 1))
        dpi = int(self.config.get("pdf_render_dpi", 190))

        images = []
        if ext == ".pdf":
            images = self.render_pdf_pages_to_images(file_path, max_pages=max_pages, dpi=dpi)
        else:
            try:
                with Image.open(file_path) as img:
                    images = [img.convert("RGB")]
            except Exception as e:
                self.log(f"Ошибка открытия изображения: {e}")
                return {}

        if not images:
            return {}

        BAD_TYPES = {"утверждение", "утверждено" "соглашение", "утверждаю", "согласование", "согласовано",
                     "подпись", "подписи", "печать", "штамп"}

        best = {}
        raw_texts = []
        header_failed = False

        for page_idx, img in enumerate(images):
            if not self.is_running:
                break

            variants = []

            if not header_failed and self.config.get("vision_try_header_first", True):
                variants.append(("header", self.build_header_crop(img, ratio=float(self.config.get("header_crop_ratio", 0.35)))))

            variants.append(("full", img))

            for variant_name, current_img in variants:
                try:
                    self.log(f"Vision: страница {page_idx + 1}, вариант {variant_name}")
                    meta = self.call_vision_model(current_img)
                    meta = self.normalize_meta_dict(meta)

                    recognized_type = meta.get("type", "").lower()
                    if recognized_type in BAD_TYPES:
                        self.log(f"⚠️ Обнаружен плохой тип '{meta.get('type')}' в варианте {variant_name}")

                        if variant_name == "header":
                            header_failed = True
                            self.log(f"🔄 Переключаюсь на полное изображение (header определил '{meta.get('type')}')")
                            continue
                        else:
                            self.log(f"⚠️ Full вариант тоже дал '{meta.get('type')}', сохраняю как есть")

                    if meta.get("raw_text"):
                        raw_texts.append(meta["raw_text"])

                    if self.is_meta_better(meta, best):
                        best = meta

                    if self.meta_is_good_enough(best):
                        break

                except requests.exceptions.ConnectionError:
                    self.log("Vision-модель недоступна: Ollama не отвечает")
                    return {}
                except requests.exceptions.Timeout:
                    self.log("Vision-модель: таймаут")
                    return {}
                except Exception as e:
                    self.log(f"Vision-модель ошибка: {e}")
                    return {}

            if self.meta_is_good_enough(best):
                break

        if raw_texts:
            merged = "\n".join([x for x in raw_texts if x.strip()]).strip()
            if merged:
                best["raw_text"] = merged

        return best

    def normalize_type(self, value: str) -> str:
        value = compact_spaces(value)
        if not value:
            return ""

        aliases = {
            "решение": "Решение",
            "приказ": "Приказ",
            "договор": "Договор",
            "акт": "Акт",
            "заявление": "Заявление",
            "счет": "Счет",
            "счёт": "Счет",
            "накладная": "Накладная",
            "счет-фактура": "Счет-фактура",
            "доверенность": "Доверенность",
            "протокол": "Протокол",
            "письмо": "Письмо",
        }

        low = value.lower()
        for key, normalized in aliases.items():
            if key in low or low == key:
                return normalized

        return value.capitalize()

    def match_type_with_templates(self, recognized_type: str) -> str:
        if not recognized_type:
            return ""

        recognized_lower = recognized_type.lower()
        available_types = self.available_doc_types()

        for tpl_type in available_types:
            if recognized_lower == tpl_type.lower():
                return tpl_type

        for tpl_type in available_types:
            if recognized_lower in tpl_type.lower() or tpl_type.lower() in recognized_lower:
                self.log(f"Сопоставление типов: '{recognized_type}' -> '{tpl_type}' (частичное совпадение)")
                return tpl_type

        keywords = {
            "решение": "Решение",
            "приказ": "Приказ",
            "договор": "Договор",
            "акт": "Акт",
            "заявление": "Заявление",
            "счет": "Счет",
            "накладная": "Накладная",
        }

        for keyword, mapped_type in keywords.items():
            if keyword in recognized_lower:
                for tpl_type in available_types:
                    if tpl_type.lower() == mapped_type.lower():
                        self.log(f"Сопоставление по ключевому слову: '{recognized_type}' -> '{tpl_type}'")
                        return tpl_type

        return ""

    def normalize_number(self, number: str) -> str:
        number = compact_spaces(number)
        if not number:
            return ""

        number = number.replace("№", "").replace("#", "")
        number = re.sub(r"(?i)\b(решение|приказ|договор|акт|заявление|номер|no|n)\b", "", number)
        number = compact_spaces(number).strip(" .,:;—-")

        if number.lower() in {"б/н", "без номера", "бн", "нет"}:
            return ""

        table = str.maketrans({
            "О": "0", "о": "0", "O": "0", "o": "0",
            "З": "3", "з": "3",
            "А": "A", "В": "B", "Е": "E", "К": "K", "М": "M",
            "Н": "H", "Р": "P", "С": "C", "Т": "T", "Х": "X",
        })
        number = number.translate(table)

        number = re.sub(r"[^0-9A-Za-zА-Яа-яЁё/\-:._]", "", number)
        number = re.sub(r"(?i)(от|дата|г|год)$", "", number).strip(" .,:;—-")

        if not re.search(r"\d", number):
            return ""
        if len(number) > 40:
            number = number[:40]
        return number

    def normalize_date(self, date_str: str) -> str:
        date_str = compact_spaces(date_str)
        if not date_str:
            return ""

        try:
            dt = datetime.strptime(date_str, "%Y-%m-%d")
            if 1900 <= dt.year <= 2100:
                return dt.strftime("%Y-%m-%d")
        except Exception:
            pass

        m = re.search(r"(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{2,4})", date_str)
        if m:
            d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
            if y < 100:
                y += 2000 if y < 50 else 1900
            try:
                return datetime(y, mo, d).strftime("%Y-%m-%d")
            except Exception:
                pass

        months = {
            "января": 1, "февраля": 2, "марта": 3, "апреля": 4, "мая": 5, "июня": 6,
            "июля": 7, "августа": 8, "сентября": 9, "октября": 10, "ноября": 11, "декабря": 12,
            "январь": 1, "февраль": 2, "март": 3, "апрель": 4, "май": 5, "июнь": 6,
            "июль": 7, "август": 8, "сентябрь": 9, "октябрь": 10, "ноябрь": 11, "декабрь": 12,
        }
        m2 = re.search(
            r"(\d{1,2})\s+(января|февраля|марта|апреля|мая|июня|июля|августа|сентября|октября|ноября|декабря|"
            r"январь|февраль|март|апрель|май|июнь|июль|август|сентябрь|октябрь|ноябрь|декабрь)\s+(\d{4})",
            date_str, re.IGNORECASE
        )
        if m2:
            d = int(m2.group(1))
            mo = months[m2.group(2).lower()]
            y = int(m2.group(3))
            try:
                return datetime(y, mo, d).strftime("%Y-%m-%d")
            except Exception:
                pass

        return ""

    def normalize_designation(self, value: str) -> str:
        value = compact_spaces(value).strip(" .,:;—-")
        if not value:
            return "Без темы"
        if len(value) > 80:
            value = " ".join(value.split()[:7])
        return value or "Без темы"

    def normalize_meta_dict(self, meta: dict) -> dict:
        if not isinstance(meta, dict):
            meta = {}
        return {
            "type": self.normalize_type(meta.get("type", "")),
            "number": self.normalize_number(meta.get("number", "")),
            "date": self.normalize_date(meta.get("date", "")),
            "designation": self.normalize_designation(meta.get("designation", "")),
            "raw_text": str(meta.get("raw_text", "") or "").strip(),
            "source": str(meta.get("source", "") or "")
        }

    def regex_extract_type(self, text: str) -> str:
        low = (text or "").lower()
        for doc_type in self.available_doc_types():
            if re.search(rf"\b{re.escape(doc_type.lower())}\b", low):
                return doc_type

        keywords = {
            "решение": "Решение",
            "приказ": "Приказ",
            "договор": "Договор",
            "акт": "Акт",
            "заявление": "Заявление",
            "счет": "Счет",
            "накладная": "Накладная",
        }
        for keyword, mapped_type in keywords.items():
            if keyword in low:
                return mapped_type
        return ""

    def regex_extract_number(self, text: str) -> str:
        if not text:
            return ""

        marker = r"(?:№|N|No|Nо|N0|Ne|Номер|номер)"
        patterns = [
            rf"(?:РЕШЕНИЕ|ПРИКАЗ|ДОГОВОР|АКТ|ЗАЯВЛЕНИЕ)\s*{marker}?\s*([0-9A-Za-zА-Яа-яЁё][0-9A-Za-zА-Яа-яЁё/\-:._]{{0,35}})",
            rf"{marker}\s*([0-9A-Za-zА-Яа-яЁё][0-9A-Za-zА-Яа-яЁё/\-:._]{{0,35}})",
            rf"\d{{1,2}}[.\-/]\d{{1,2}}[.\-/]\d{{2,4}}\s*{marker}?\s*([0-9A-Za-zА-Яа-яЁё][0-9A-Za-zА-Яа-яЁё/\-:._]{{0,35}})"
        ]

        bad = {"от", "г", "год", "решение", "приказ", "договор", "акт", "заявление", "дата"}

        for pattern in patterns:
            for m in re.finditer(pattern, text, flags=re.IGNORECASE):
                raw = m.group(1).strip(" .,:;—-")
                if raw.lower() in bad:
                    continue
                n = self.normalize_number(raw)
                if n:
                    return n
        return ""

    def regex_extract_date(self, text: str) -> str:
        if not text:
            return ""

        candidates = []
        for m in re.finditer(r"\b(\d{1,2}[.\-/]\d{1,2}[.\-/]\d{2,4})\b", text):
            candidates.append(m.group(1))
            if len(candidates) >= 10:
                break

        month_pattern = (
            r"\b\d{1,2}\s+"
            r"(?:января|февраля|марта|апреля|мая|июня|июля|августа|сентября|октября|ноября|декабря)"
            r"\s+\d{4}\b"
        )
        for m in re.finditer(month_pattern, text, flags=re.IGNORECASE):
            candidates.append(m.group(0))
            if len(candidates) >= 15:
                break

        for c in candidates:
            d = self.normalize_date(c)
            if d:
                return d
        return ""

    def apply_regex_fallbacks(self, meta: dict, text: str) -> dict:
        meta = self.normalize_meta_dict(meta)

        if not meta.get("type"):
            meta["type"] = self.regex_extract_type(text)

        if not meta.get("number"):
            meta["number"] = self.regex_extract_number(text)

        if not meta.get("date"):
            meta["date"] = self.regex_extract_date(text)

        if not meta.get("designation") or meta.get("designation") == "Без темы":
            lines = [compact_spaces(x) for x in text.splitlines() if compact_spaces(x)]
            for line in lines[:20]:
                if 8 <= len(line) <= 90:
                    if not re.search(r"(?i)^\s*(№|N|No|\d{1,2}[.\-/]\d{1,2})", line):
                        if not re.search(r"(?i)\b(решение|приказ|договор|акт)\b", line):
                            meta["designation"] = self.normalize_designation(line)
                            break

        return self.normalize_meta_dict(meta)

    def meta_score(self, meta: dict) -> int:
        meta = self.normalize_meta_dict(meta)
        score = 0
        if meta.get("type"):
            score += 35
        if meta.get("number"):
            score += 25
        if meta.get("date"):
            score += 25
        if meta.get("designation") and meta.get("designation") != "Без темы":
            score += 10
        if len(meta.get("raw_text", "")) > 80:
            score += 5
        return score

    def is_meta_better(self, new: dict, old: dict) -> bool:
        return self.meta_score(new) > self.meta_score(old)

    def meta_is_good_enough(self, meta: dict) -> bool:
        meta = self.normalize_meta_dict(meta)
        return bool(meta.get("type") and (meta.get("number") or meta.get("date")))

    def ddmmyyyy(self, date_iso: str) -> str:
        try:
            return datetime.strptime(date_iso, "%Y-%m-%d").strftime("%d.%m.%Y")
        except Exception:
            return ""

    def select_best_template(self, meta: dict, templates: list[dict]) -> dict:
        if not templates:
            return None

        if len(templates) == 1:
            return templates[0]

        self.log(f"Выбор лучшего шаблона из {len(templates)} вариантов")

        number = meta.get("number", "")
        date_str = meta.get("date", "")
        designation = meta.get("designation", "")
        raw_text = meta.get("raw_text", "")

        best_template = templates[0]
        best_score = -1

        for template in templates:
            score = 0
            filename_tpl = template.get("filename", "")
            subdir = template.get("subdir", "")

            if "{номер}" in filename_tpl and number:
                score += 3
            elif "{номер}" not in filename_tpl and not number:
                score += 2

            if "{дата}" in filename_tpl and date_str:
                score += 3
            elif "{дата}" not in filename_tpl and not date_str:
                score += 2

            if "{тема}" in filename_tpl and designation != "Без темы":
                score += 2

            if subdir.lower() in raw_text.lower():
                score += 2

            self.log(f"  Шаблон '{subdir}': score={score}")

            if score > best_score:
                best_score = score
                best_template = template

        self.log(f"Выбран шаблон: '{best_template.get('subdir', '')}' (score={best_score})")
        return best_template

    def extract_meta_from_file(self, file_path: str) -> tuple[dict, str]:
        ext = Path(file_path).suffix.lower()
        raw_text = ""

        vision_meta = {}
        if self.config.get("vision_enabled", True):
            vision_meta = self.call_vision_for_file(file_path)
            vision_meta["source"] = "vision"
            raw_text = vision_meta.get("raw_text", "") or ""
            vision_meta = self.apply_regex_fallbacks(vision_meta, raw_text)

            if self.meta_is_good_enough(vision_meta):
                self.log(f"Vision модель OK: type={vision_meta.get('type')}, score={self.meta_score(vision_meta)}")
                return vision_meta, raw_text
            else:
                self.log(f"Vision модель неполный результат: type={vision_meta.get('type')}, score={self.meta_score(vision_meta)}")

        if not raw_text.strip():
            self.log("Vision не распознал текст документа")
            return vision_meta, raw_text

        return vision_meta, raw_text

    def find_templates_by_type(self, doc_type: str) -> list[dict]:
        doc_type = compact_spaces(doc_type).lower()
        out = []
        for tpl in self.templates.get("templates", []):
            if compact_spaces(tpl.get("doc_type", "")).lower() == doc_type:
                out.append(tpl)
        return out

    def build_new_name(self, meta: dict, template: dict, original_ext: str) -> tuple[str, str]:
        number = self.normalize_number(meta.get("number", ""))
        date_iso = self.normalize_date(meta.get("date", ""))
        designation = self.normalize_designation(meta.get("designation", ""))

        values = {
            "тип": template.get("doc_type", "Документ"),
            "номер": number if number else "без номера",
            "дата": self.ddmmyyyy(date_iso) if date_iso else "без даты",
            "тема": designation or template.get("doc_type", "Документ"),
        }

        filename_tpl = template.get("filename", "{тип} №{номер} от {дата}")
        try:
            name = filename_tpl.format(**values)
        except Exception:
            name = "{тип} №{номер} от {дата}".format(**values)

        name = re.sub(r"№\s*№+", "№", name)
        name = self.sanitize_filename(name)

        subdir = self.sanitize_filename(template.get("subdir", template.get("doc_type", "")))
        return subdir, name

    def build_unknown_name(self, meta: dict, original_ext: str) -> str:
        number = self.normalize_number(meta.get("number", ""))
        date_iso = self.normalize_date(meta.get("date", ""))

        parts = ["Неизвестный"]

        if number:
            parts.append(f"№{number}")

        if date_iso:
            date_formatted = self.ddmmyyyy(date_iso)
            if date_formatted:
                parts.append(f"от {date_formatted}")

        name = " ".join(parts)
        return self.sanitize_filename(name)

    def move_to_unknown(self, src_path: str, reason: str, meta: dict | None = None):
        unknown_dir = self.config.get("unknown_dir", "unknown")
        ensure_dir(unknown_dir)

        original_name = Path(src_path).name
        base = self.sanitize_filename(Path(src_path).stem)
        ext = Path(src_path).suffix.lower()

        if meta:
            base = self.build_unknown_name(meta, ext)

        dest = os.path.join(unknown_dir, f"{base}{ext}")
        dest = self.get_unique_path(dest)

        new_name = Path(dest).name

        shutil.move(src_path, dest)
        self.log(f"[UNKNOWN] {original_name} -> {new_name} (причина: {reason})")

        self.results['unknown'] += 1
        self.results['details'].append({
            'filename': new_name,
            'original_name': original_name,
            'type': meta.get('type', 'Неизвестный') if meta else 'Неизвестный',
            'status': 'UNKNOWN',
            'destination': f"unknown/{new_name}",
            'reason': reason
        })

        self.write_history({
            "src": str(src_path),
            "dest": str(dest),
            "status": "unknown",
            "reason": reason,
            "meta": meta or {},
            "timestamp": datetime.now().isoformat()
        })

    def process_file(self, file_path: str):
        ensure_dir(self.config.get("out_dir", "sorted"))
        ensure_dir(self.config.get("unknown_dir", "unknown"))
        ensure_dir("logs")

        original_name = Path(file_path).name

        size_limit_mb = int(self.config.get("max_file_size_mb", 80))
        size_mb = os.path.getsize(file_path) / (1024 * 1024)
        if size_mb > size_limit_mb:
            self.move_to_unknown(file_path, f"too_big_{int(size_mb)}mb")
            return

        meta, raw_text = self.extract_meta_from_file(file_path)
        meta = self.apply_regex_fallbacks(meta, raw_text)

        Path("logs/last_ocr_text.txt").write_text(raw_text[:30000], encoding="utf-8")
        Path("logs/last_meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")

        if not raw_text.strip():
            self.move_to_unknown(file_path, "no_text", meta)
            return

        recognized_type_raw = meta.get("type", "")
        recognized_type_normalized = self.normalize_type(recognized_type_raw)

        BAD_TYPES = {"утверждение", "утверждено" "соглашение", "утверждаю", "согласование", "согласовано",
                     "подпись", "подписи", "печать", "штамп"}
        if recognized_type_normalized.lower() in BAD_TYPES:
            self.log(f"⚠️ Итоговый тип '{recognized_type_normalized}' считается плохим, отправляю в unknown")
            self.move_to_unknown(file_path, f"bad_type_{recognized_type_normalized}", meta)
            return

        self.log(f"Распознанный тип: '{recognized_type_raw}' -> нормализованный: '{recognized_type_normalized}'")

        matched_type = self.match_type_with_templates(recognized_type_normalized)

        if not matched_type:
            self.move_to_unknown(file_path, f"type_mismatch", meta)
            return

        self.log(f"Тип сопоставлен с шаблоном: '{matched_type}'")
        meta["type"] = matched_type

        matching_templates = self.find_templates_by_type(matched_type)
        if not matching_templates:
            self.move_to_unknown(file_path, f"no_template_for_type_{matched_type}", meta)
            return

        best_template = self.select_best_template(meta, matching_templates)
        if not best_template:
            self.move_to_unknown(file_path, f"no_suitable_template", meta)
            return

        ext = Path(file_path).suffix.lower()
        subdir, name_no_ext = self.build_new_name(meta, best_template, ext)
        dest_dir = os.path.join(self.config.get("out_dir", "sorted"), subdir) if subdir else self.config.get("out_dir", "sorted")
        ensure_dir(dest_dir)

        dest_path = os.path.join(dest_dir, name_no_ext + ext)
        dest_path = self.get_unique_path(dest_path)

        shutil.move(file_path, dest_path)
        self.log(f"[OK] {original_name} -> {name_no_ext + ext} (шаблон: {best_template.get('subdir', '')})")

        self.results['processed'] += 1
        try:
            rel_path = os.path.relpath(dest_path, os.getcwd())
        except ValueError:
            rel_path = dest_path

        self.results['details'].append({
            'filename': name_no_ext + ext,
            'original_name': original_name,
            'type': matched_type,
            'status': 'OK',
            'destination': rel_path
        })

        self.write_history({
            "src": str(file_path),
            "dest": str(dest_path),
            "status": "ok",
            "meta": meta,
            "template": best_template,
            "score": self.meta_score(meta),
            "timestamp": datetime.now().isoformat()
        })

    def run(self):
        try:
            if not os.path.exists(self.inbox_dir):
                self.log(f"Папка входящих не существует: {self.inbox_dir}")
                self.signals.finished.emit(self.results)
                return

            if self.config.get("vision_enabled", True):
                self.check_ollama_available(self.config.get("vision_url", "http://localhost:11434/api/generate"), "Ollama Vision")

            files = [
                str(p) for p in Path(self.inbox_dir).iterdir()
                if p.is_file() and p.suffix.lower() in self.SUPPORTED_EXTENSIONS
            ]

            total = len(files)
            self.results['total'] = total

            if total == 0:
                self.log(f"Нет файлов. Поддерживаются: {', '.join(sorted(self.SUPPORTED_EXTENSIONS))}")
                self.signals.finished.emit(self.results)
                return

            self.log(f"Найдено файлов: {total}")
            self.log(f"Доступные типы документов из шаблонов: {', '.join(self.available_doc_types())}")
            self.log(f"Vision={'да' if self.config.get('vision_enabled', True) else 'нет'}")

            for i, file_path in enumerate(files):
                if not self.is_running:
                    self.log("Обработка остановлена пользователем")
                    break

                name = Path(file_path).name
                self.signals.progress.emit(int((i + 1) / total * 100), f"Обработка: {name}")
                self.log(f"Обработка: {name}")

                try:
                    self.process_file(file_path)
                except Exception as e:
                    self.log(f"[ERROR] {name}: {e}")
                    self.log(traceback.format_exc())
                    try:
                        self.move_to_unknown(file_path, f"error_{type(e).__name__}")
                    except Exception:
                        pass

            self.signals.progress.emit(100, "Завершено")
            self.log(f"Обработка завершена. Успешно: {self.results['processed']}, неизвестных: {self.results['unknown']}")
        except Exception as e:
            self.signals.error.emit(str(e))
        finally:
            self.signals.finished.emit(self.results)

    def check_ollama_available(self, url: str, title: str) -> bool:
        try:
            tags_url = url.replace("/api/generate", "/api/tags")
            r = requests.get(tags_url, timeout=4)
            if r.status_code == 200:
                models = r.json().get("models", [])
                self.log(f"{title} доступен. Моделей: {len(models)}")
                return True
        except Exception:
            pass
        self.log(f"{title} недоступен")
        return False

    def stop(self):
        self.is_running = False


# =========================
# GUI: Настройки
# =========================

class ConfigWidget(QWidget):
    def __init__(self, config: dict):
        super().__init__()
        self.config = config.copy()
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout()

        paths_group = QGroupBox("Основные настройки")
        paths_layout = QVBoxLayout()

        self.inbox_edit = self.add_path_row(paths_layout, "Папка входящих:", self.config.get("inbox_dir", "scannerdata"))
        self.out_edit = self.add_path_row(paths_layout, "Папка результатов:", self.config.get("out_dir", "sorted"))
        self.unknown_edit = self.add_path_row(paths_layout, "Папка неизвестных:", self.config.get("unknown_dir", "unknown"))

        interval_layout = QHBoxLayout()
        interval_layout.addWidget(QLabel("Интервал проверки (минут):"))
        self.interval_spin = QSpinBox()
        self.interval_spin.setRange(1, 1440)
        self.interval_spin.setValue(int(self.config.get("check_interval_minutes", 5)))
        interval_layout.addWidget(self.interval_spin)
        paths_layout.addLayout(interval_layout)

        size_layout = QHBoxLayout()
        size_layout.addWidget(QLabel("Макс. размер файла, MB:"))
        self.max_file_size_spin = QSpinBox()
        self.max_file_size_spin.setRange(1, 1000)
        self.max_file_size_spin.setValue(int(self.config.get("max_file_size_mb", 80)))
        size_layout.addWidget(self.max_file_size_spin)
        paths_layout.addLayout(size_layout)

        self.notify_check = QCheckBox("Показывать уведомления о результатах")
        self.notify_check.setChecked(bool(self.config.get("show_notifications", True)))
        paths_layout.addWidget(self.notify_check)

        self.show_logs_check = QCheckBox("Показывать вкладку с логами")
        self.show_logs_check.setChecked(bool(self.config.get("show_logs_tab", True)))
        self.show_logs_check.stateChanged.connect(self.toggle_logs_tab)
        paths_layout.addWidget(self.show_logs_check)

        paths_group.setLayout(paths_layout)
        layout.addWidget(paths_group)

        vision_group = QGroupBox("Vision-модель")
        vision_layout = QVBoxLayout()

        self.vision_enabled_check = QCheckBox("Использовать vision-модель")
        self.vision_enabled_check.setChecked(bool(self.config.get("vision_enabled", True)))
        vision_layout.addWidget(self.vision_enabled_check)

        vision_layout.addWidget(QLabel("Vision URL:"))
        self.vision_url_edit = QLineEdit(self.config.get("vision_url", "http://localhost:11434/api/generate"))
        vision_layout.addWidget(self.vision_url_edit)

        vision_layout.addWidget(QLabel("Vision model:"))
        self.vision_model_edit = QLineEdit(self.config.get("vision_model", "qwen2.5vl:7b"))
        vision_layout.addWidget(self.vision_model_edit)

        vision_params = QHBoxLayout()
        vision_params.addWidget(QLabel("PDF DPI:"))
        self.pdf_dpi_spin = QSpinBox()
        self.pdf_dpi_spin.setRange(120, 300)
        self.pdf_dpi_spin.setValue(int(self.config.get("pdf_render_dpi", 190)))
        vision_params.addWidget(self.pdf_dpi_spin)

        vision_params.addWidget(QLabel("Страниц:"))
        self.vision_pages_spin = QSpinBox()
        self.vision_pages_spin.setRange(1, 5)
        self.vision_pages_spin.setValue(int(self.config.get("vision_max_pages", 1)))
        vision_params.addWidget(self.vision_pages_spin)

        vision_params.addWidget(QLabel("Таймаут Vision (0 = без таймаута):"))
        self.vision_timeout_spin = QSpinBox()
        self.vision_timeout_spin.setRange(0, 9999)
        self.vision_timeout_spin.setValue(int(self.config.get("vision_timeout", 0)))
        vision_params.addWidget(self.vision_timeout_spin)

        vision_layout.addLayout(vision_params)

        self.header_first_check = QCheckBox("Сначала анализировать верхнюю часть страницы")
        self.header_first_check.setChecked(bool(self.config.get("vision_try_header_first", True)))
        vision_layout.addWidget(self.header_first_check)

        header_layout = QHBoxLayout()
        header_layout.addWidget(QLabel("Доля верхней части:"))
        self.header_ratio_spin = QSpinBox()
        self.header_ratio_spin.setRange(20, 60)
        self.header_ratio_spin.setValue(int(float(self.config.get("header_crop_ratio", 0.35)) * 100))
        header_layout.addWidget(self.header_ratio_spin)
        header_layout.addWidget(QLabel("%"))
        vision_layout.addLayout(header_layout)

        vision_layout.addWidget(QLabel("Доп. подсказка для vision-модели:"))
        self.vision_user_prompt_edit = QTextEdit()
        self.vision_user_prompt_edit.setPlainText(self.config.get("vision_user_prompt", ""))
        self.vision_user_prompt_edit.setPlaceholderText("Например: номер обычно находится после слова 'Решение №', а дата в верхней части документа.")
        vision_layout.addWidget(self.vision_user_prompt_edit)

        vision_group.setLayout(vision_layout)
        layout.addWidget(vision_group)

        layout.addStretch()
        self.setLayout(layout)

    def add_path_row(self, parent_layout: QVBoxLayout, label: str, value: str) -> QLineEdit:
        row = QHBoxLayout()
        row.addWidget(QLabel(label))
        edit = QLineEdit(value)
        row.addWidget(edit)
        btn = QPushButton("Обзор...")
        btn.clicked.connect(lambda: self.select_dir(edit))
        row.addWidget(btn)
        parent_layout.addLayout(row)
        return edit

    def select_dir(self, edit: QLineEdit):
        path = QFileDialog.getExistingDirectory(self, "Выберите папку")
        if path:
            edit.setText(path)

    def toggle_logs_tab(self):
        parent = self.parent()
        while parent is not None:
            if isinstance(parent, MainWindow):
                parent.toggle_logs_tab(self.show_logs_check.isChecked())
                break
            parent = parent.parent()

    def get_config(self) -> dict:
        return {
            "inbox_dir": self.inbox_edit.text().strip() or "scannerdata",
            "out_dir": self.out_edit.text().strip() or "sorted",
            "unknown_dir": self.unknown_edit.text().strip() or "unknown",
            "check_interval_minutes": self.interval_spin.value(),
            "max_file_size_mb": self.max_file_size_spin.value(),
            "show_notifications": self.notify_check.isChecked(),
            "show_logs_tab": self.show_logs_check.isChecked(),

            "vision_enabled": self.vision_enabled_check.isChecked(),
            "vision_url": self.vision_url_edit.text().strip(),
            "vision_model": self.vision_model_edit.text().strip() or "qwen2.5vl:7b",
            "vision_timeout": self.vision_timeout_spin.value(),
            "vision_max_pages": self.vision_pages_spin.value(),
            "vision_num_ctx": 8192,
            "vision_try_header_first": self.header_first_check.isChecked(),
            "header_crop_ratio": self.header_ratio_spin.value() / 100.0,
            "vision_user_prompt": self.vision_user_prompt_edit.toPlainText().strip(),

            "pdf_render_dpi": self.pdf_dpi_spin.value(),
        }


# =========================
# GUI: Шаблоны
# =========================

class TemplatesWidget(QWidget):
    def __init__(self, templates: dict):
        super().__init__()
        self.templates = templates.get("templates", [])
        if not isinstance(self.templates, list):
            self.templates = []
        self.current_index = None
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout()

        buttons = QHBoxLayout()
        self.add_btn = QPushButton("+ Добавить")
        self.add_btn.clicked.connect(self.add_template)
        self.delete_btn = QPushButton("- Удалить")
        self.delete_btn.clicked.connect(self.delete_template)
        buttons.addWidget(self.add_btn)
        buttons.addWidget(self.delete_btn)
        buttons.addStretch()
        layout.addLayout(buttons)

        self.template_list = QListWidget()
        self.template_list.currentRowChanged.connect(self.on_template_selected)
        layout.addWidget(self.template_list)

        group = QGroupBox("Редактор шаблона")
        group_layout = QVBoxLayout()

        group_layout.addWidget(QLabel("Тип документа:"))
        self.doc_type_edit = QLineEdit()
        group_layout.addWidget(self.doc_type_edit)

        group_layout.addWidget(QLabel("Шаблон имени файла:"))
        self.filename_edit = QLineEdit()
        group_layout.addWidget(self.filename_edit)

        info = QLabel("Переменные: {тип}, {номер}, {дата}, {тема}")
        info.setStyleSheet("QLabel { color: #666; }")
        group_layout.addWidget(info)

        group_layout.addWidget(QLabel("Подпапка:"))
        self.subdir_edit = QLineEdit()
        group_layout.addWidget(self.subdir_edit)

        self.update_btn = QPushButton("Сохранить изменения")
        self.update_btn.clicked.connect(self.update_template)
        group_layout.addWidget(self.update_btn)

        group.setLayout(group_layout)
        layout.addWidget(group)

        self.setLayout(layout)
        self.refresh_list()

    def refresh_list(self):
        selected = self.current_index
        self.template_list.blockSignals(True)
        self.template_list.clear()

        counts = {}
        for tpl in self.templates:
            doc_type = tpl.get("doc_type", "")
            counts[doc_type] = counts.get(doc_type, 0) + 1

        seen = {}
        for tpl in self.templates:
            doc_type = tpl.get("doc_type", "")
            subdir = tpl.get("subdir", "")
            seen[doc_type] = seen.get(doc_type, 0) + 1
            suffix = f" (вариант {seen[doc_type]})" if counts.get(doc_type, 0) > 1 else ""
            self.template_list.addItem(f"{doc_type} → {subdir}{suffix}")

        self.template_list.blockSignals(False)

        if selected is not None and 0 <= selected < len(self.templates):
            self.template_list.setCurrentRow(selected)
        elif self.templates:
            self.template_list.setCurrentRow(0)
        else:
            self.clear_editor()

    def clear_editor(self):
        self.current_index = None
        self.doc_type_edit.clear()
        self.filename_edit.clear()
        self.subdir_edit.clear()

    def on_template_selected(self, index: int):
        if index < 0 or index >= len(self.templates):
            self.clear_editor()
            return
        self.current_index = index
        tpl = self.templates[index]
        self.doc_type_edit.setText(tpl.get("doc_type", ""))
        self.filename_edit.setText(tpl.get("filename", ""))
        self.subdir_edit.setText(tpl.get("subdir", ""))

    def add_template(self):
        self.templates.append({
            "doc_type": "Новый тип",
            "filename": "{тип} №{номер} от {дата}",
            "subdir": "Новая папка"
        })
        self.current_index = len(self.templates) - 1
        self.refresh_list()

    def delete_template(self):
        idx = self.template_list.currentRow()
        if idx < 0 or idx >= len(self.templates):
            QMessageBox.warning(self, "Ошибка", "Выберите шаблон")
            return

        reply = QMessageBox.question(
            self,
            "Удаление",
            f"Удалить шаблон '{self.templates[idx].get('doc_type', '')}'?",
            QMessageBox.Yes | QMessageBox.No
        )
        if reply == QMessageBox.Yes:
            del self.templates[idx]
            self.current_index = min(idx, len(self.templates) - 1) if self.templates else None
            self.refresh_list()

    def update_template(self):
        idx = self.current_index
        if idx is None or idx < 0 or idx >= len(self.templates):
            QMessageBox.warning(self, "Ошибка", "Выберите шаблон")
            return

        doc_type = self.doc_type_edit.text().strip()
        filename = self.filename_edit.text().strip()
        subdir = self.subdir_edit.text().strip()

        if not doc_type:
            QMessageBox.warning(self, "Ошибка", "Тип документа не может быть пустым")
            return
        if not filename:
            QMessageBox.warning(self, "Ошибка", "Шаблон имени файла не может быть пустым")
            return

        self.templates[idx] = {
            "doc_type": doc_type,
            "filename": filename,
            "subdir": subdir
        }
        self.current_index = idx
        self.refresh_list()
        QMessageBox.information(self, "Готово", "Шаблон сохранён")

    def get_templates(self) -> dict:
        return {"templates": self.templates}


# =========================
# Вкладка "Главная"
# =========================

class HomeWidget(QWidget):
    def __init__(self, start_callback):
        super().__init__()
        self.start_callback = start_callback
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout()
        layout.setAlignment(Qt.AlignCenter)

        title = QLabel("ДокСорт")
        title.setStyleSheet("""
            font-size: 36pt; 
            font-weight: bold; 
            color: #2c3e50;
            padding: 20px;
        """)
        title.setAlignment(Qt.AlignCenter)
        layout.addWidget(title)

        subtitle = QLabel("Автоматическая сортировка и распознавание документов")
        subtitle.setStyleSheet("""
            font-size: 14pt; 
            color: #7f8c8d;
            padding: 10px;
        """)
        subtitle.setAlignment(Qt.AlignCenter)
        layout.addWidget(subtitle)

        layout.addSpacing(50)

        self.start_btn = QPushButton("СТАРТ")
        self.start_btn.setStyleSheet("""
            QPushButton {
                font-size: 24pt; 
                font-weight: bold;
                padding: 30px 80px;
                background: qlineargradient(x1:0, y1:0, x2:1, y2:1,
                    stop:0 #4CAF50, stop:1 #45a049);
                color: white;
                border: none;
                border-radius: 15px;
                min-width: 300px;
            }
            QPushButton:hover {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:1,
                    stop:0 #66BB6A, stop:1 #4CAF50);
                transform: scale(1.05);
            }
            QPushButton:pressed {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:1,
                    stop:0 #388E3C, stop:1 #2E7D32);
            }
        """)
        self.start_btn.clicked.connect(self.start_callback)
        layout.addWidget(self.start_btn, alignment=Qt.AlignCenter)

        layout.addSpacing(30)

        self.setLayout(layout)


# =========================
# Вкладка "Об авторе"
# =========================

class AboutWidget(QWidget):
    def __init__(self):
        super().__init__()
        self.init_ui()

    def init_ui(self):
        layout = QVBoxLayout()
        layout.setAlignment(Qt.AlignCenter)

        title = QLabel("Об авторе")
        title.setStyleSheet("""
            font-size: 24pt; 
            font-weight: bold; 
            color: #2c3e50;
            padding: 20px;
        """)
        title.setAlignment(Qt.AlignCenter)
        layout.addWidget(title)

        layout.addSpacing(30)

        info_label = QLabel(
            "Разработчик данного ПО: Рашитов Артур\n\n"
            "Для сотрудничества обращаться по почте:\n"
            "rashar03@mail.ru"
        )
        info_label.setStyleSheet("""
            font-size: 14pt; 
            color: #34495e;
            padding: 20px;
            line-height: 1.5;
        """)
        info_label.setAlignment(Qt.AlignCenter)
        layout.addWidget(info_label)

        layout.addSpacing(50)

        version_label = QLabel("ДокСорт v2.3 (15.05.2026)")
        version_label.setStyleSheet("font-size: 10pt; color: #95a5a6;")
        version_label.setAlignment(Qt.AlignCenter)
        layout.addWidget(version_label)

        self.setLayout(layout)


# =========================
# Главное окно
# =========================

class MainWindow(QMainWindow):
    def __init__(self, single_instance: SingleInstance = None):
        super().__init__()
        self.single_instance = single_instance
        self.worker = None
        self.timer = None
        self.is_processing = False
        self.load_configs()
        self.init_ui()
        self.init_tray()

        # Запускаем сервер для приема команд от других экземпляров
        if self.single_instance:
            self.single_instance.start_server(self.show_from_tray)

    def show_from_tray(self):
        """Показывает окно из трея."""
        self.show()
        self.raise_()
        self.activateWindow()

    def default_config(self) -> dict:
        return {
            "inbox_dir": "scannerdata",
            "out_dir": "sorted",
            "unknown_dir": "unknown",
            "check_interval_minutes": 5,
            "max_file_size_mb": 80,
            "show_notifications": True,
            "show_logs_tab": False,

            "vision_enabled": True,
            "vision_url": "http://localhost:11434/api/generate",
            "vision_model": "qwen2.5vl:7b",
            "vision_timeout": 0,
            "vision_max_pages": 1,
            "vision_num_ctx": 8192,
            "vision_try_header_first": True,
            "header_crop_ratio": 0.35,
            "vision_user_prompt": "",

            "pdf_render_dpi": 190,
        }

    def default_templates(self) -> dict:
        return {
            "templates": [
                {
                    "doc_type": "Решение",
                    "filename": "Решение №{номер} от {дата}",
                    "subdir": "Решения",
                },
                {
                    "doc_type": "Приказ",
                    "filename": "Приказ №{номер} от {дата}",
                    "subdir": "Приказы",
                },
                {
                    "doc_type": "Договор",
                    "filename": "Договор №{номер} от {дата}",
                    "subdir": "Договоры",
                },
                {
                    "doc_type": "Акт",
                    "filename": "Акт №{номер} от {дата}",
                    "subdir": "Акты",
                },
                {
                    "doc_type": "Заявление",
                    "filename": "Заявление {тема} от {дата}",
                    "subdir": "Заявления",
                },
            ]
        }

    def load_configs(self):
        base = self.default_config()
        loaded = safe_yaml_load("config.yaml", base)
        base.update(loaded)
        self.config = base

        self.templates = safe_yaml_load("templates.yaml", self.default_templates())
        if "templates" not in self.templates or not isinstance(self.templates["templates"], list):
            self.templates = self.default_templates()

    def init_ui(self):
        self.setWindowTitle("ДокСорт - Умный сортировщик документов")
        self.setGeometry(100, 100, 1000, 700)

        icon_path = resource_path("icon.ico")
        if os.path.exists(icon_path):
            self.setWindowIcon(QIcon(icon_path))
        else:
            try:
                if getattr(sys, 'frozen', False):
                    base_path = sys._MEIPASS
                    temp_icon = os.path.join(base_path, "icon.ico")
                    if os.path.exists(temp_icon):
                        self.setWindowIcon(QIcon(temp_icon))
            except:
                pass

        central = QWidget()
        self.setCentralWidget(central)
        main_layout = QVBoxLayout()
        central.setLayout(main_layout)

        self.tabs = QTabWidget()

        self.home_widget = HomeWidget(self.toggle_processing)

        self.config_widget = ConfigWidget(self.config)
        self.templates_widget = TemplatesWidget(self.templates)
        self.about_widget = AboutWidget()

        self.log_widget = QWidget()
        log_layout = QVBoxLayout()
        self.log_text = QTextEdit()
        self.log_text.setReadOnly(True)
        self.log_text.setFont(QFont("Consolas", 9))
        log_layout.addWidget(self.log_text)
        clear_btn = QPushButton("Очистить лог")
        clear_btn.clicked.connect(self.log_text.clear)
        log_layout.addWidget(clear_btn)
        self.log_widget.setLayout(log_layout)

        self.tabs.addTab(self.home_widget, "Главная")
        self.tabs.addTab(self.templates_widget, "Шаблоны")
        self.tabs.addTab(self.config_widget, "Настройки")
        self.log_index = self.tabs.addTab(self.log_widget, "Логи")
        self.tabs.addTab(self.about_widget, "Об авторе")

        if not self.config.get("show_logs_tab", True):
            self.tabs.setTabVisible(self.log_index, False)

        main_layout.addWidget(self.tabs)

        bottom = QHBoxLayout()

        self.status_label = QLabel("Готов к работе")
        bottom.addWidget(self.status_label)

        bottom.addStretch()

        save_btn = QPushButton("Сохранить настройки")
        save_btn.clicked.connect(self.save_configs)
        bottom.addWidget(save_btn)

        main_layout.addLayout(bottom)

    def toggle_logs_tab(self, visible: bool):
        """Показывает или скрывает вкладку с логами"""
        if hasattr(self, 'log_index'):
            self.tabs.setTabVisible(self.log_index, visible)

    def create_icon_with_dot(self, base_icon_path: str, has_dot: bool = False) -> QIcon:
        icon_path = resource_path(base_icon_path)

        if not os.path.exists(icon_path):
            return self.style().standardIcon(self.style().SP_ComputerIcon)

        base_pixmap = QPixmap(icon_path)

        if not has_dot:
            return QIcon(base_pixmap)

        scaled_pixmap = base_pixmap.scaled(64, 64, Qt.KeepAspectRatio, Qt.SmoothTransformation)

        from PyQt5.QtGui import QPainter
        painter = QPainter(scaled_pixmap)

        dot_radius = 12
        margin = 4
        x = scaled_pixmap.width() - dot_radius - margin
        y = margin + dot_radius

        painter.setBrush(QBrush(QColor(255, 0, 0)))
        painter.setPen(Qt.NoPen)
        painter.drawEllipse(x - dot_radius, y - dot_radius, dot_radius * 2, dot_radius * 2)

        painter.setBrush(Qt.NoBrush)
        painter.setPen(QColor(255, 255, 255, 180))
        painter.drawEllipse(x - dot_radius, y - dot_radius, dot_radius * 2, dot_radius * 2)

        painter.end()

        return QIcon(scaled_pixmap)

    def init_tray(self):
        self.tray_icon = QSystemTrayIcon(self)
        icon = self.create_icon_with_dot("icon.ico", has_dot=False)
        self.tray_icon.setIcon(icon)

        self.tray_icon.setToolTip("ДокСорт")

        tray_menu = QMenu()

        show_action = QAction("Показать", self)
        show_action.triggered.connect(self.show)
        tray_menu.addAction(show_action)

        start_stop_action = QAction("Старт/Стоп", self)
        start_stop_action.triggered.connect(self.toggle_processing)
        tray_menu.addAction(start_stop_action)

        tray_menu.addSeparator()

        quit_action = QAction("Выход", self)
        quit_action.triggered.connect(self.quit_app)
        tray_menu.addAction(quit_action)

        self.tray_icon.setContextMenu(tray_menu)
        self.tray_icon.activated.connect(self.tray_activated)
        self.tray_icon.show()

    def tray_activated(self, reason):
        if reason == QSystemTrayIcon.DoubleClick:
            self.show()
            self.raise_()

    def closeEvent(self, event):
        event.ignore()
        self.hide()
        self.tray_icon.showMessage(
            "ДокСорт",
            "Приложение свернуто в трей. Обработка продолжится в фоне.",
            QSystemTrayIcon.Information,
            2000
        )

    def update_tray_icon(self):
        icon = self.create_icon_with_dot("icon.ico", has_dot=self.is_processing)
        self.tray_icon.setIcon(icon)

        if self.is_processing:
            self.tray_icon.setToolTip("ДокСорт - ИДЕТ ОБРАБОТКА")
        else:
            self.tray_icon.setToolTip("ДокСорт")

    def update_start_button(self):
        """Обновляет текст кнопки на главной вкладке"""
        if self.is_processing:
            self.home_widget.start_btn.setText("СТОП")
            self.home_widget.start_btn.setStyleSheet("""
                QPushButton {
                    font-size: 24pt; 
                    font-weight: bold;
                    padding: 30px 80px;
                    background: qlineargradient(x1:0, y1:0, x2:1, y2:1,
                        stop:0 #f44336, stop:1 #da190b);
                    color: white;
                    border: none;
                    border-radius: 15px;
                    min-width: 300px;
                }
                QPushButton:hover {
                    background: qlineargradient(x1:0, y1:0, x2:1, y2:1,
                        stop:0 #ef5350, stop:1 #f44336);
                }
                QPushButton:pressed {
                    background: qlineargradient(x1:0, y1:0, x2:1, y2:1,
                        stop:0 #c62828, stop:1 #b71c1c);
                }
            """)
        else:
            self.home_widget.start_btn.setText("СТАРТ")
            self.home_widget.start_btn.setStyleSheet("""
                QPushButton {
                    font-size: 24pt; 
                    font-weight: bold;
                    padding: 30px 80px;
                    background: qlineargradient(x1:0, y1:0, x2:1, y2:1,
                        stop:0 #4CAF50, stop:1 #45a049);
                    color: white;
                    border: none;
                    border-radius: 15px;
                    min-width: 300px;
                }
                QPushButton:hover {
                    background: qlineargradient(x1:0, y1:0, x2:1, y2:1,
                        stop:0 #66BB6A, stop:1 #4CAF50);
                }
                QPushButton:pressed {
                    background: qlineargradient(x1:0, y1:0, x2:1, y2:1,
                        stop:0 #388E3C, stop:1 #2E7D32);
                }
            """)

    def validate_processing_start(self) -> tuple:
        """
        Проверяет доступность Ollama перед запуском обработки.
        Возвращает (успех, сообщение_об_ошибке)
        """
        # Проверяем доступность vision-модели если она включена
        if self.config.get("vision_enabled", True):
            vision_url = self.config.get("vision_url", "http://localhost:11434/api/generate")

            # Проверяем, похож ли URL на Ollama
            if "11434" in vision_url or "ollama" in vision_url.lower():
                if not self.check_ollama_availability():
                    return False, self.get_ollama_error_message()
            else:
                # Проверяем общую доступность эндпоинта
                if not self.check_vision_endpoint(vision_url):
                    return False, f"Vision-модель недоступна по адресу:\n{vision_url}\n\nПроверьте:\n1. Запущен ли сервер vision-модели\n2. Правильно ли указан URL в настройках\n3. Доступен ли порт"

        return True, ""

    def check_ollama_availability(self) -> bool:
        """Проверяет доступность Ollama и наличие нужной модели."""
        try:
            vision_url = self.config.get("vision_url", "http://localhost:11434/api/generate")
            vision_model = self.config.get("vision_model", "qwen2.5vl:7b")

            # Проверяем доступность API
            tags_url = vision_url.replace("/api/generate", "/api/tags")
            r = requests.get(tags_url, timeout=5)

            if r.status_code != 200:
                return False

            # Проверяем наличие модели
            models = r.json().get("models", [])
            model_names = [m.get("name", "") for m in models]

            # Ищем точное совпадение или модель с префиксом
            model_found = any(
                vision_model == name or name.startswith(vision_model.split(":")[0])
                for name in model_names
            )

            return model_found

        except requests.exceptions.ConnectionError:
            return False
        except Exception:
            return False

    def get_ollama_error_message(self) -> str:
        """Формирует подробное сообщение об ошибке Ollama."""
        vision_url = self.config.get("vision_url", "http://localhost:11434/api/generate")
        vision_model = self.config.get("vision_model", "qwen2.5vl:7b")

        # Проверяем разные аспекты проблемы
        try:
            tags_url = vision_url.replace("/api/generate", "/api/tags")
            r = requests.get(tags_url, timeout=3)

            if r.status_code == 200:
                # Сервер доступен, но модели нет
                models = r.json().get("models", [])
                model_names = [m.get("name", "") for m in models]

                if not model_names:
                    return (
                        f"Ollama запущена, но нет ни одной модели.\n\n"
                    )
                else:
                    return (
                            f"Модель '{vision_model}' не найдена в Ollama.\n\n"
                            f"Доступные модели:\n" +
                            "\n".join(f"  • {name}" for name in model_names[:10]) +
                            f"\n\nУстановите нужную модель.\n"

                    )
        except:
            pass

        # Сервер недоступен
        return (
            f"Ollama не запущена или недоступна.\n\n"
            f"Проверьте:\n"
            f"1. Запущена ли Ollama\n"
            f"2. Правильно ли указан URL: {vision_url}\n"
            f"3. Доступен ли порт (не заблокирован ли файерволом)\n\n"
        )

    def check_vision_endpoint(self, url: str) -> bool:
        """Проверяет доступность любого vision-эндпоинта."""
        try:
            r = requests.get(url, timeout=5)
            return r.status_code < 500
        except requests.exceptions.ConnectionError:
            return False
        except requests.exceptions.Timeout:
            return False
        except Exception:
            return False

    def toggle_processing(self):
        if self.is_processing:
            self.stop_processing()
        else:
            # Проверяем возможность запуска
            self.config = self.config_widget.get_config()
            self.templates = self.templates_widget.get_templates()

            can_start, error_message = self.validate_processing_start()

            if not can_start:
                # Показываем подробное сообщение об ошибке
                msg_box = QMessageBox(self)
                msg_box.setIcon(QMessageBox.Warning)
                msg_box.setWindowTitle("Невозможно запустить обработку")
                msg_box.setText("Не удалось запустить обработку документов")
                msg_box.setInformativeText(error_message)
                msg_box.setDetailedText(error_message)

                # Добавляем кастомные кнопки
                settings_btn = msg_box.addButton("Открыть настройки", QMessageBox.ActionRole)
                close_btn = msg_box.addButton(QMessageBox.Ok)

                msg_box.exec_()

                # Если нажали "Открыть настройки"
                if msg_box.clickedButton() == settings_btn:
                    self.tabs.setCurrentWidget(self.config_widget)

                self.update_start_button()
                return

            # Все проверки пройдены, запускаем обработку
            self.save_configs(silent=True)
            self.save_templates(silent=True)
            self.start_processing_impl()

        self.update_start_button()

    def start_processing_impl(self):
        """Внутренний метод для запуска обработки после проверок."""
        inbox = self.config.get("inbox_dir", "scannerdata")

        self.is_processing = True
        self.update_tray_icon()
        self.status_label.setText("Работаю...")

        interval = self.config.get("check_interval_minutes", 5) * 60 * 1000
        self.timer = QTimer()
        self.timer.timeout.connect(self.run_worker)
        self.timer.start(interval)

        self.run_worker()

        self.hide()
        self.tray_icon.showMessage(
            "ДокСорт",
            f"Автоматическая обработка запущена.\nПроверка каждые {self.config.get('check_interval_minutes', 5)} мин.",
            QSystemTrayIcon.Information,
            3000
        )

    def stop_processing(self):
        self.is_processing = False
        self.update_tray_icon()
        if self.timer:
            self.timer.stop()
            self.timer = None

        if self.worker:
            self.worker.stop()
            self.worker = None

        self.status_label.setText("Остановлен")

        self.tray_icon.showMessage(
            "ДокСорт",
            f"Автоматическая обработка остановлена.",
            QSystemTrayIcon.Information,
            3000
        )

    def run_worker(self):
        if self.worker and self.worker.is_alive():
            return

        inbox = self.config.get("inbox_dir", "scannerdata")

        self.worker = DocumentProcessorWorker(self.config, self.templates, inbox)
        self.worker.signals.finished.connect(self.on_processing_finished)
        self.worker.signals.error.connect(self.on_processing_error)
        self.worker.signals.log.connect(self.on_processing_log)
        self.worker.start()
        if not os.path.exists(inbox):
            reply = QMessageBox.question(
                self,
                "Папка не найдена",
                f"Папка входящих не существует:\n{inbox}\n\nСоздать?",
                QMessageBox.Yes | QMessageBox.No
            )
            if reply == QMessageBox.Yes:
                os.makedirs(inbox, exist_ok=True)
            else:
                return

    def on_processing_finished(self, results):
        self.worker = None

        if results and results.get('total', 0) > 0:
            self.log_message(f"Обработка завершена. Успешно: {results.get('processed', 0)}, неизвестных: {results.get('unknown', 0)}")

            if self.config.get("show_notifications", True):
                dialog = ResultsDialog(results, self.templates, self.config.get("inbox_dir", "scannerdata"), self)
                dialog.exec_()

                self.templates = dialog.templates
                self.templates_widget.templates = self.templates.get("templates", [])
                self.templates_widget.refresh_list()
        else:
            self.log_message("Нет новых файлов для обработки")

    def on_processing_error(self, msg: str):
        """Обрабатывает ошибки обработки."""
        self.log_message(f"ОШИБКА: {msg}")

        # Дополнительно сохраняем ошибку в отдельный лог
        try:
            ensure_dir("logs")
            error_entry = {
                "timestamp": now_ts(),
                "error": msg,
                "type": "error"
            }
            with open("logs/errors.jsonl", "a", encoding="utf-8") as f:
                f.write(json.dumps(error_entry, ensure_ascii=False) + "\n")
        except Exception:
            pass

        QMessageBox.critical(self, "Ошибка", msg)

    def on_processing_log(self, msg: str):
        self.log_message(msg)

    def log_message(self, msg: str):
        """Выводит сообщение в лог и сохраняет в файл."""
        timestamp = now_ts()
        formatted_msg = f"[{timestamp}] {msg}"

        # Выводим в текстовое поле
        self.log_text.append(formatted_msg)
        self.log_text.verticalScrollBar().setValue(
            self.log_text.verticalScrollBar().maximum()
        )

        # Сохраняем в файл логов в том же формате
        try:
            ensure_dir("logs")
            with open("logs/history.log", "a", encoding="utf-8") as f:
                f.write(formatted_msg + "\n")
        except Exception:
            pass  # Игнорируем ошибки записи логов

    # Сохраняем в файл логов
    try:
        ensure_dir("logs")
        log_entry = {
            "timestamp": timestamp,
            "message": msg,
            "type": "log"
        }
        with open("logs/history.jsonl", "a", encoding="utf-8") as f:
            f.write(json.dumps(log_entry, ensure_ascii=False) + "\n")
    except Exception:
        pass  # Игнорируем ошибки записи логов

    def save_configs(self, silent: bool = False):
        cfg = self.config_widget.get_config()
        with open("config.yaml", "w", encoding="utf-8") as f:
            yaml.dump(cfg, f, allow_unicode=True, default_flow_style=False, sort_keys=False)
        if not silent:
            QMessageBox.information(self, "Готово", "Настройки сохранены в config.yaml")

    def save_templates(self, silent: bool = False):
        tpl = self.templates_widget.get_templates()
        with open("templates.yaml", "w", encoding="utf-8") as f:
            yaml.dump(tpl, f, allow_unicode=True, default_flow_style=False, sort_keys=False)
        if not silent:
            QMessageBox.information(self, "Готово", "Шаблоны сохранены в templates.yaml")

    def quit_app(self):
        if self.is_processing:
            self.stop_processing()
        if self.single_instance:
            self.single_instance.cleanup()
        self.tray_icon.hide()
        QApplication.quit()


def main():
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    app.setQuitOnLastWindowClosed(False)

    # Проверяем, запущен ли уже экземпляр
    single_instance = SingleInstance()

    if single_instance.try_connect():
        # Другой экземпляр уже запущен - отправляем сигнал и выходим
        print("Приложение уже запущено. Активируем существующий экземпляр.")
        sys.exit(0)

    # Создаем главное окно с поддержкой SingleInstance
    window = MainWindow(single_instance)
    window.show()

    sys.exit(app.exec_())


if __name__ == "__main__":
    main()